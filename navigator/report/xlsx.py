"""The Excel workbook: every table of the report, numbers kept as numbers (decision 0014-5)."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .data import ReportData, Table

FORMATS = {"eur": '#,##0 "€"', "pct": "0.0%", "num": "#,##0", "num1": "#,##0.0", "int": "0"}
HEAD = PatternFill("solid", fgColor="F4F1EA")


def _sheet_name(title: str, used: set[str]) -> str:
    bad = '[]:*?/\\'
    name = "".join("-" if ch in bad else ch for ch in title)[:31] or "Sheet"
    base, k = name, 2
    while name in used:
        name = f"{base[:28]}_{k}"
        k += 1
    used.add(name)
    return name


def _write(ws, t: Table, row: int = 1) -> int:
    ws.cell(row, 1, t.title).font = Font(bold=True, size=12, color="0F6E78")
    row += 1
    if t.note:
        ws.cell(row, 1, t.note).alignment = Alignment(wrap_text=False)
        row += 1
    for c, h in enumerate(t.header, start=1):
        cell = ws.cell(row, c, h)
        cell.font = Font(bold=True, color="56677A")
        cell.fill = HEAD
    row += 1
    fmts = t.formats or ["text"] * len(t.header)
    for r in t.rows:
        for c, v in enumerate(r, start=1):
            cell = ws.cell(row, c, v)
            f = fmts[c - 1] if c - 1 < len(fmts) else "text"
            if f in FORMATS and isinstance(v, (int, float)):
                cell.number_format = FORMATS[f]
        row += 1
    for c in range(1, len(t.header) + 1):
        width = max([len(str(t.header[c - 1]))] + [len(str(r[c - 1])) for r in t.rows if c - 1 < len(r)] + [8])
        ws.column_dimensions[get_column_letter(c)].width = min(width + 2, 60)
    return row + 1


def build_xlsx(data: ReportData, path: str | Path) -> Path:
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    used: set[str] = set()
    ws.title = _sheet_name(data.labels.get("summary", "Summary"), used)
    ws.cell(1, 1, data.title).font = Font(bold=True, size=14)
    ws.cell(2, 1, data.subtitle)
    ws.cell(3, 1, f"{data.date} · Decarbonisation Navigator v{data.version}")
    r = 5
    for line in data.summary:
        ws.cell(r, 1, line)
        r += 1
    r += 1
    for k, v in data.settings:
        ws.cell(r, 1, k).font = Font(color="56677A")
        ws.cell(r, 2, v)
        r += 1
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 40
    tables = [t for t in (data.position.get("table_full"), data.measures, data.plans.get("table"),
                          data.plans.get("adaptive", {}).get("table") if data.plans.get("adaptive") else None,
                          data.weights.get("table"), data.ranking, data.smaa.get("table"), data.assumptions)
              if t is not None] + list(data.tables)
    for t in tables:
        _write(wb.create_sheet(_sheet_name(t.title, used)), t)
    if data.sources:
        src = Table(data.labels.get("sources", "Sources"), ["Title", "URL", "Checked"], [list(s) for s in data.sources])
        _write(wb.create_sheet(_sheet_name(src.title, used)), src)
    wb.save(path)
    return path
